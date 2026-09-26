"""Real PostgreSQL/pgvector tests; only external storage and model calls are faked."""
from __future__ import annotations

import io
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
import pytest_asyncio
from docx import Document
from fastapi import HTTPException, UploadFile
from sqlalchemy import make_url, text
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.api.auth import Principal, RequestContext
from app.api.v1 import gap_audits, marketing, policies
from app.context.gap_auditor import audit_signal, reduce_assessment, verify_criterion
from app.context.gap_models import MarketingRequest, ReviewRequest, RunRequest
from app.context.marketing_checker import check_copy
from app.context.policy_service import DOCX, chunk_sections, extract_sections, index_policy, validate_document
from app.core.config import get_settings
from app.synthesis.obligation_extractor import SourceRequired, extract_obligations, validate_source

pytestmark = [pytest.mark.integration, pytest.mark.asyncio]

CLAUSES = [
    ('1.1','Record retention','Retain customer records for five years.'),
    ('1.2','Suspicious activity reporting','Report suspicious transactions to the designated authority.'),
    ('1.3','Access review','Review access permissions every quarter.'),
]
SOURCE = '\n'.join(f'{reference} {criterion}' for reference,_,criterion in CLAUSES) + '\n' + (
    'This is an artificial integration-test circular, not an actual regulatory instrument. ' * 20)


class Model:
    model = 'test-only'
    last_model = 'test-only'
    last_provider = 'test-only'

    async def generate(self, *, context, **kwargs):
        if 'source' in context:
            return {'obligations': [dict(clause_reference=reference,requirement_title=title,
                assessment_criteria=[criterion],applicable_departments=['compliance_legal'],
                source_excerpt=criterion,statutory_sanction=None,statutory_deadline=None)
                for reference,title,criterion in CLAUSES]}
        excerpt = context['excerpts'][0]
        return {'verdict':'satisfied','evidence':[{'chunk_id':excerpt['chunk_id'],'excerpt':excerpt['content']}],
                'reasoning':'The test policy explicitly states the record retention period.'}

    async def aclose(self):
        pass


class Embedder:
    async def embed(self, inputs):
        return tuple(tuple([1.0,0.0]+[0.0]*1534) if 'records' in item.lower()
                     else tuple([0.0,1.0]+[0.0]*1534) for item in inputs)

    async def aclose(self):
        pass


@pytest_asyncio.fixture
async def ctx(monkeypatch):
    settings = get_settings()
    assert settings.ENVIRONMENT == 'test'
    url = make_url(settings.DATABASE_URL)
    assert url.host in {'localhost','127.0.0.1','::1'} and url.database == 'sc_test'
    monkeypatch.setattr(settings,'REGULATORY_GAP_ENABLED',True)
    monkeypatch.setattr(policies,'store_file',AsyncMock(return_value='test-object-version'))
    monkeypatch.setattr(policies,'dispatch',lambda *args: True)
    monkeypatch.setattr(gap_audits,'dispatch',lambda *args: True)
    engine = create_async_engine(url.set(drivername='postgresql+asyncpg'))
    org,user = uuid4(),uuid4()
    async with engine.connect() as connection:
        transaction = await connection.begin()
        try:
            await connection.execute(text("SELECT set_config('app.current_tenant_id',:org,true)"),{'org':str(org)})
            await connection.execute(text('INSERT INTO auth.tenants(id,name,slug) VALUES(:id,:name,:slug)'),
                {'id':org,'name':'Gap engine isolated test','slug':str(org)})
            await connection.execute(text("INSERT INTO auth.users(id,tenant_id,email,permission_role) VALUES(:id,:org,:email,'ADMIN')"),
                {'id':user,'org':org,'email':f'{user}@example.invalid'})
            await connection.execute(text('SET LOCAL ROLE sc_app_runtime'))
            async with AsyncSession(bind=connection,join_transaction_mode='create_savepoint',expire_on_commit=False) as session:
                yield RequestContext(Principal(user_id=user,tenant_id=org,permission_role='ADMIN',
                    permissions=frozenset({'READ_INTELLIGENCE','REVIEW_COMPLIANCE','MANAGE_POLICIES','USE_CIL'})),session)
        finally:
            await transaction.rollback()
    await engine.dispose()


async def source_signal(ctx):
    row = (await ctx.session.execute(text('''INSERT INTO pipeline.signals
        (signal_type,title,body_text,source_url,detected_at,pipeline_stage,dedup_status,is_proprietary)
        VALUES('regulatory_mandate','Artificial test circular',:body,'https://www.cbn.gov.ng/test-fixture-only',
               now(),'NORMALIZED','UNIQUE',false) RETURNING *'''),{'body':SOURCE})).mappings().one()
    return dict(row)


def policy_document():
    document = Document()
    document.add_paragraph('Retain customer records for five years.')
    table = document.add_table(rows=1,cols=1)
    table.cell(0,0).text = 'Records are retained securely for five years.'
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def ready_policy(ctx):
    body = policy_document()
    response = await policies.upload_policy(UploadFile(filename='controls.docx',file=io.BytesIO(body)),
        'Example AML policy','aml_kyc','1.0',None,ctx)
    row = (await ctx.session.execute(text('SELECT * FROM organizations.tenant_policies WHERE id=:id'),
        {'id':response['id']})).mappings().one()
    await index_policy(ctx.session,dict(row),client=Embedder(),body=body)
    return response['id']


async def assessed(ctx):
    signal = await source_signal(ctx)
    await ready_policy(ctx)
    run = await gap_audits.create_run(RunRequest(signal_id=signal['id'],idempotency_key=uuid4()),ctx)
    await audit_signal(ctx.session,run,signal,embedder=Embedder(),verifier=Model())
    return signal,await gap_audits.list_audits(signal['id'],ctx)


async def test_obligation_extraction_persists_source_and_criteria_idempotently(ctx):
    signal = await source_signal(ctx)
    extraction,rows = await extract_obligations(ctx.session,signal,client=Model())
    assert len(rows) == 3 and all(row['assessment_criteria'] for row in rows)
    again,again_rows = await extract_obligations(ctx.session,signal,client=Model())
    assert again == extraction and [row['id'] for row in rows] == [row['id'] for row in again_rows]
    stored = (await ctx.session.execute(text('SELECT source_text FROM pipeline.regulatory_extractions WHERE id=:id'),{'id':extraction})).scalar_one()
    assert stored == SOURCE


async def test_policy_upload_extracts_docx_tables_and_persists_vectors(ctx):
    identifier = await ready_policy(ctx)
    rows = (await ctx.session.execute(text('''SELECT content,vector_dims(embedding) dimensions,location
        FROM organizations.policy_chunks WHERE policy_id=:id'''),{'id':identifier})).mappings().all()
    assert rows and all(row['dimensions']==1536 for row in rows)
    assert 'Records are retained securely' in rows[0]['content']
    assert any('table' in part for part in rows[0]['location']['sources'])


async def test_gap_determination_and_live_pgvector_tenant_scoping(ctx):
    signal,result = await assessed(ctx)
    statuses = {row['clause_reference']:row['status'] for row in result['items']}
    assert statuses == {'1.1':'adequately_met','1.2':'gap_deficient','1.3':'gap_deficient'}
    assert result['run']['processing_status'] == 'completed'
    await ctx.session.execute(text("SELECT set_config('app.current_tenant_id',:org,true)"),{'org':str(uuid4())})
    assert (await ctx.session.execute(text('SELECT count(*) FROM organizations.policy_chunks'))).scalar_one() == 0
    assert (await ctx.session.execute(text('SELECT count(*) FROM pipeline.compliance_gap_audits'))).scalar_one() == 0
    assert (await ctx.session.execute(text('SELECT count(*) FROM audit.compliance_gap_events'))).scalar_one() == 0


async def test_reviewer_override_is_atomic_versioned_idempotent_and_immutable(ctx):
    _,result = await assessed(ctx)
    audit = next(row for row in result['items'] if row['status']=='gap_deficient')
    payload = ReviewRequest(status='adequately_met',reason='Reviewed supporting operational evidence.',expected_revision=audit['revision'],idempotency_key=uuid4())
    updated = await gap_audits.override(audit['id'],payload,ctx)
    assert updated['status']=='adequately_met' and updated['automated_status']=='gap_deficient'
    assert updated['revision']==audit['revision']+1
    await gap_audits.override(audit['id'],payload,ctx)
    events = (await gap_audits.history(audit['id'],ctx))['items']
    assert len(events)==2 and events[-1]['event_type']=='override'
    with pytest.raises(HTTPException) as conflict:
        await gap_audits.override(audit['id'],payload.model_copy(update={'idempotency_key':uuid4()}),ctx)
    assert conflict.value.status_code==409
    for statement in ["UPDATE audit.compliance_gap_events SET reason='tampered'",'DELETE FROM audit.compliance_gap_events','TRUNCATE audit.compliance_gap_events']:
        with pytest.raises(Exception):
            async with ctx.session.begin_nested():
                await ctx.session.execute(text(statement))


async def test_marketing_checker_cites_claims_and_persists_result(ctx):
    result = await marketing.check_campaign(MarketingRequest(copy='Earn guaranteed 25% returns. Bank with us. Send money worldwide without limits.',channel='sms'),ctx)
    codes = {finding['rule_id'] for finding in result['findings']}
    assert {'investment_promises','banking_scope','remittance_scope','opt_out'} <= codes
    assert all(item['source_url'].startswith('https://') for item in result['findings'])
    assert result['approved'] is False
    assert (await ctx.session.execute(text('SELECT count(*) FROM pipeline.marketing_checks WHERE id=:id'),{'id':result['id']})).scalar_one()==1


async def test_reviewer_permission_and_cross_tenant_id_access(ctx):
    from dataclasses import replace
    _,result = await assessed(ctx)
    audit = result['items'][0]
    payload = ReviewRequest(reason='Reviewed all available documents.',expected_revision=1,idempotency_key=uuid4())
    viewer = RequestContext(replace(ctx.principal,permissions=frozenset({'READ_INTELLIGENCE'})),ctx.session)
    with pytest.raises(HTTPException) as denied:
        await gap_audits.sign_off(audit['id'],payload,viewer)
    assert denied.value.status_code==403
    other = uuid4()
    await ctx.session.execute(text("SELECT set_config('app.current_tenant_id',:org,true)"),{'org':str(other)})
    foreign = RequestContext(replace(ctx.principal,tenant_id=other),ctx.session)
    with pytest.raises(HTTPException) as denied:
        await gap_audits.sign_off(audit['id'],payload,foreign)
    assert denied.value.status_code==404


async def test_provider_failure_does_not_turn_into_a_compliance_gap(ctx):
    class Broken(Model):
        async def generate(self,**kwargs):
            raise TimeoutError('test provider unavailable')
    signal = await source_signal(ctx)
    with pytest.raises(TimeoutError):
        await extract_obligations(ctx.session,signal,client=Broken())
    assert (await ctx.session.execute(text('SELECT count(*) FROM pipeline.regulatory_obligations WHERE signal_id=:id'),{'id':signal['id']})).scalar_one()==0


async def test_chunking_and_document_validation():
    chunks = chunk_sections([('customer records retention ' * 800,{'page':1})])
    assert len(chunks)>1 and all(chunk.token_count<=500 for chunk in chunks)
    assert chunks[0].token_count==500
    assert validate_document(policy_document(),'policy.docx')==DOCX
    assert extract_sections(policy_document(),DOCX)
    with pytest.raises(ValueError):
        validate_document(b'not a PDF','policy.pdf')
    with pytest.raises(SourceRequired):
        validate_source({'body_text':'Short summary','source_url':'https://cbn.gov.ng/circular','executive_summary':'Short summary'})


async def test_similarity_boundaries_and_verifier_evidence_validation():
    candidate = {'id':uuid4(),'policy_id':uuid4(),'content':'Retain records for five years.',
        'document_title':'Retention','version':'1','content_sha256':'a'*64,'location':{},'similarity_score':.82}
    result = await verify_criterion('Retain records',[candidate],Model())
    assert result['verdict']=='partial'
    candidate['similarity_score']=.83
    assert (await verify_criterion('Retain records',[candidate],Model()))['verdict']=='satisfied'
    candidate['similarity_score']=.64
    assert (await verify_criterion('Retain records',[candidate],Model()))['verdict']=='missing'
    assert reduce_assessment([{'verdict':'satisfied'},{'verdict':'missing'}])==('partially_met',50.0)
    assert reduce_assessment([{'verdict':'satisfied'},{'verdict':'contradicted'}])[0]=='gap_deficient'
    assert not check_copy('Returns are not guaranteed. Capital is at risk.','social',[])['findings']


async def test_addendum_signoff_and_changed_policy_versions(ctx):
    _,result = await assessed(ctx)
    audit = result['items'][0]
    policy = (await ctx.session.execute(text('SELECT id FROM organizations.tenant_policies WHERE active'))).scalar_one()
    updated = await gap_audits.addendum(audit['id'],ReviewRequest(reason='Supporting approved policy document attached.',
        expected_revision=audit['revision'],idempotency_key=uuid4(),policy_id=policy),ctx)
    await gap_audits.sign_off(audit['id'],ReviewRequest(reason='Reviewed this exact assessment and its evidence.',
        expected_revision=updated['revision'],idempotency_key=uuid4()),ctx)
    events = (await gap_audits.history(audit['id'],ctx))['items']
    assert [event['event_type'] for event in events] == ['assessed','addendum','sign_off']
    await ready_policy(ctx)
    with pytest.raises(HTTPException) as stale:
        await gap_audits.sign_off(audit['id'],ReviewRequest(reason='Attempted signoff after the policy set changed.',
            expected_revision=updated['revision'],idempotency_key=uuid4()),ctx)
    assert stale.value.status_code==409


async def test_pdf_text_locations_and_scanned_pdf_handling():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    from app.context.policy_service import PDF
    writer = PdfWriter()
    page = writer.add_blank_page(width=612,height=792)
    font = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (Retain customer records for five years.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    body = io.BytesIO()
    writer.write(body)
    sections = extract_sections(body.getvalue(),PDF)
    assert 'Retain customer records' in sections[0][0] and sections[0][1]['page']==1
    blank = PdfWriter()
    blank.add_blank_page(width=612,height=792)
    body = io.BytesIO()
    blank.write(body)
    assert extract_sections(body.getvalue(),PDF)==[]


async def test_expired_trial_cannot_use_evidence_endpoints(ctx):
    await ctx.session.execute(text("UPDATE auth.tenants SET pilot_expires_at=now()-interval '1 day' WHERE id=:org"),{'org':ctx.principal.tenant_id})
    with pytest.raises(HTTPException) as expired:
        await policies.list_policies(ctx)
    assert expired.value.status_code==402


async def test_empty_extraction_requests_source_review_without_inventing_obligations(ctx):
    class AbstainingModel(Model):
        async def generate(self, **kwargs):
            return {'obligations': []}
    signal = await source_signal(ctx)
    with pytest.raises(SourceRequired, match='INSUFFICIENT_BINDING_OBLIGATIONS'):
        await extract_obligations(ctx.session, signal, client=AbstainingModel())
    assert (await ctx.session.execute(text('SELECT count(*) FROM pipeline.regulatory_obligations WHERE signal_id=:id'),
        {'id': signal['id']})).scalar_one() == 0


async def test_pending_run_keeps_completed_health_and_blocks_stale_review(ctx):
    signal, result = await assessed(ctx)
    before = (await policies.list_policies(ctx))['health']
    assert before['assessed_obligations'] == 3 and before['stale_obligations'] == 0
    await gap_audits.create_run(RunRequest(signal_id=signal['id'], idempotency_key=uuid4()), ctx)
    assert (await policies.list_policies(ctx))['health'] == before
    audit = result['items'][0]
    for action in ('override', 'addendum', 'sign_off'):
        payload = ReviewRequest(reason='Reviewer attempts to change an earlier assessment.',
            expected_revision=audit['revision'], idempotency_key=uuid4(),
            status='adequately_met' if action == 'override' else None)
        with pytest.raises(HTTPException) as stale:
            await gap_audits.review(audit['id'], payload, action, ctx)
        assert stale.value.status_code == 409
    await ready_policy(ctx)
    health = (await policies.list_policies(ctx))['health']
    assert health['stale_obligations'] == 3 and health['assessed_obligations'] == 3


async def test_late_old_job_cannot_overwrite_a_newer_completed_audit(ctx, monkeypatch):
    from app.workers.tasks import regulatory_gap
    signal, result = await assessed(ctx)
    older = result['run']['id']
    newer = await gap_audits.create_run(RunRequest(signal_id=signal['id'], idempotency_key=uuid4()), ctx)
    await audit_signal(ctx.session, newer, signal, embedder=Embedder(), verifier=Model())
    await ctx.session.execute(text("UPDATE pipeline.compliance_gap_runs SET processing_status='queued' WHERE id=:id"), {'id': older})
    await ctx.session.commit()
    async def sessions():
        yield ctx.session
    monkeypatch.setattr(regulatory_gap, 'get_session', sessions)
    result = await regulatory_gap.run_job('audit', str(ctx.principal.tenant_id), str(older))
    assert result['status'] == 'superseded'
    assert set((await ctx.session.execute(text('SELECT run_id FROM pipeline.compliance_gap_audits'))).scalars()) == {newer['id']}


@pytest.mark.parametrize('enqueue_fails', [False, True])
async def test_policy_activation_and_dependent_runs_are_atomic(ctx, monkeypatch, enqueue_fails):
    from app.workers.tasks import regulatory_gap
    signal, _ = await assessed(ctx)
    body = policy_document()
    uploaded = await policies.upload_policy(UploadFile(filename='new.docx', file=io.BytesIO(body)),
        'New controls', 'aml_kyc', '1.0', None, ctx)
    async def sessions():
        yield ctx.session
    async def index(session, policy):
        return await index_policy(session, policy, client=Embedder(), body=body)
    monkeypatch.setattr(regulatory_gap, 'get_session', sessions)
    monkeypatch.setattr(regulatory_gap, 'index_policy', index)
    monkeypatch.setattr(regulatory_gap, 'dispatch', lambda *args: False)
    if enqueue_fails:
        monkeypatch.setattr(regulatory_gap, 'enqueue_audit', AsyncMock(side_effect=RuntimeError('enqueue unavailable')))
    result = await regulatory_gap.run_job('policy', str(ctx.principal.tenant_id), str(uploaded['id']))
    policy = (await ctx.session.execute(text('SELECT active,processing_status FROM organizations.tenant_policies WHERE id=:id'),
        {'id': uploaded['id']})).mappings().one()
    runs = (await ctx.session.execute(text('SELECT count(*) FROM pipeline.compliance_gap_runs WHERE signal_id=:signal'),
        {'signal': signal['id']})).scalar_one()
    assert policy['active'] is (not enqueue_fails)
    assert policy['processing_status'] == ('failed' if enqueue_fails else 'ready')
    assert result['status'] == ('failed' if enqueue_fails else 'completed')
    assert runs == (1 if enqueue_fails else 2)


@pytest.mark.parametrize('copy', [
    'Earn guaranteed 25.5% returns.',
    'Earn guaranteed 25.5%\nreturns.',
    '🚀 Earn guaranteed\n25% returns. Terms apply.',
])
async def test_marketing_decimal_and_multiline_claims_preserve_offsets(copy):
    findings = check_copy(copy, 'social', [])['findings']
    assert any(item['rule_id'] == 'investment_promises' for item in findings)
    assert all(copy[item['start']:item['end']] == item['excerpt'] for item in findings)


async def test_legacy_relative_circular_is_resolved_using_its_incoming_feed(ctx, monkeypatch):
    from app.synthesis import obligation_extractor
    incoming = uuid4()
    relative = '/Out/2026/CCD/fixture.pdf'
    await ctx.session.execute(text('''INSERT INTO pipeline.incoming_signals
        (id,source_name,source_url,content_hash,status) VALUES(:id,'CBN Circulars',:url,:hash,'promoted')'''),
        {'id': incoming, 'url': relative, 'hash': str(uuid4())})
    signal = await source_signal(ctx)
    signal.update({'incoming_signal_id': incoming, 'source_url': relative})
    async def hydrate(value):
        assert value['source_url'] == 'https://www.cbn.gov.ng' + relative
        return value
    monkeypatch.setattr(obligation_extractor, 'hydrate_source', hydrate)
    _, obligations = await extract_obligations(ctx.session, signal, client=Model())
    assert len(obligations) == 3
