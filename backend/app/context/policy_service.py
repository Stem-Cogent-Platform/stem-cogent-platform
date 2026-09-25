"""Bounded PDF/DOCX ingestion with versioned tenant evidence and token chunks."""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass

import boto3
import tiktoken
from docx import Document
from pypdf import PdfReader
from sqlalchemy import text

from app.core.config import get_settings
from app.core.secrets import get_scalar_secret
from app.intelligence.embeddings.client import OpenAIEmbeddingClient

PDF = 'application/pdf'
DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
MODEL = 'text-embedding-3-small'


class PolicyInputError(ValueError):
    pass


@dataclass(frozen=True)
class Chunk:
    content: str
    token_count: int
    location: dict


def embedding_client() -> OpenAIEmbeddingClient:
    settings = get_settings()
    key = get_scalar_secret(settings.OPENAI_API_KEY_ARN) if settings.OPENAI_API_KEY_ARN else settings.OPENAI_API_KEY
    if not key:
        raise RuntimeError('EMBEDDING_NOT_CONFIGURED')
    return OpenAIEmbeddingClient(key, MODEL, 1536,
                                 timeout_seconds=settings.EMBEDDING_TIMEOUT_SECONDS,
                                 max_retries=settings.EMBEDDING_MAX_RETRIES)


def validate_document(body: bytes, filename: str) -> str:
    if not body or len(body) > min(get_settings().POLICY_MAX_UPLOAD_BYTES, 20971520):
        raise PolicyInputError('Document must be between 1 byte and 20 MB')
    if filename.lower().endswith('.pdf') and body.startswith(b'%PDF-'):
        return PDF
    if filename.lower().endswith('.docx') and body.startswith(b'PK'):
        try:
            with zipfile.ZipFile(io.BytesIO(body)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(m.file_size for m in members) > 100 * 1024 * 1024:
                    raise PolicyInputError('DOCX expanded content is too large')
                if 'word/document.xml' not in archive.namelist():
                    raise PolicyInputError('Invalid DOCX document')
        except zipfile.BadZipFile as exc:
            raise PolicyInputError('Invalid DOCX document') from exc
        return DOCX
    raise PolicyInputError('Upload a valid PDF or DOCX document')


def extract_sections(body: bytes, content_type: str) -> list[tuple[str, dict]]:
    if content_type == PDF:
        reader = PdfReader(io.BytesIO(body), strict=False)
        if reader.is_encrypted:
            raise PolicyInputError('Encrypted PDFs must be unlocked before upload')
        if len(reader.pages) > 500:
            raise PolicyInputError('Documents are limited to 500 pages')
        sections = [(page.extract_text() or '', {'page': index + 1})
                    for index, page in enumerate(reader.pages)]
    elif content_type == DOCX:
        document = Document(io.BytesIO(body))
        sections = [(paragraph.text, {'paragraph': index + 1})
                    for index, paragraph in enumerate(document.paragraphs)]
        sections.extend(('\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows),
                         {'table': index + 1}) for index, table in enumerate(document.tables))
    else:
        raise PolicyInputError('Unsupported document type')
    if sum(len(value) for value, _ in sections) > 2_000_000:
        raise PolicyInputError('Extracted document is too large')
    return [(value.strip(), location) for value, location in sections if value.strip()]


def chunk_sections(sections: list[tuple[str, dict]]) -> list[Chunk]:
    """500 tokens, 100-token overlap, retaining source locations across boundaries."""
    encoding = tiktoken.get_encoding('cl100k_base')
    tokens, spans = [], []
    for content, location in sections:
        start = len(tokens)
        tokens.extend(encoding.encode(content + '\n', disallowed_special=()))
        spans.append((start, len(tokens), location))
    chunks = []
    for start in range(0, len(tokens), 400):
        end = min(start + 500, len(tokens))
        content = encoding.decode(tokens[start:end], errors='ignore').strip()
        if content:
            locations = [location for left, right, location in spans if left < end and right > start]
            chunks.append(Chunk(content, end - start, {'sources': locations}))
        if end == len(tokens):
            break
    if len(chunks) > 2000:
        raise PolicyInputError('Document exceeds 2000 chunks')
    return chunks


async def store_file(body: bytes, key: str, content_type: str) -> str | None:
    settings = get_settings()
    if not settings.S3_ENTERPRISE_UPLOADS_BUCKET:
        raise RuntimeError('POLICY_STORAGE_NOT_CONFIGURED')
    client = boto3.client('s3', region_name=settings.AWS_REGION)
    response = await asyncio.to_thread(client.put_object,
        Bucket=settings.S3_ENTERPRISE_UPLOADS_BUCKET, Key=key, Body=body,
        ContentType=content_type, Metadata={'sha256': hashlib.sha256(body).hexdigest()})
    return response.get('VersionId')


async def read_file(policy: dict) -> bytes:
    settings = get_settings()
    client = boto3.client('s3', region_name=settings.AWS_REGION)
    request = {'Bucket': settings.S3_ENTERPRISE_UPLOADS_BUCKET, 'Key': policy['file_s3_key']}
    if policy.get('file_s3_version'):
        request['VersionId'] = policy['file_s3_version']
    response = await asyncio.to_thread(client.get_object, **request)
    stream = response['Body']
    try:
        body = await asyncio.to_thread(stream.read, 20971521)
    finally:
        stream.close()
    if hashlib.sha256(body).hexdigest() != policy['content_sha256']:
        raise PolicyInputError('Stored document checksum mismatch')
    return body


async def index_policy(session, policy: dict, *, client=None, body: bytes | None = None) -> int:
    body = body if body is not None else await read_file(policy)
    sections = await asyncio.to_thread(extract_sections, body, policy['content_type'])
    chunks = await asyncio.to_thread(chunk_sections, sections)
    if not chunks:
        raise PolicyInputError('NEEDS_OCR')
    owned = client is None
    client = client or embedding_client()
    try:
        for offset in range(0, len(chunks), 32):
            batch = chunks[offset:offset + 32]
            vectors = await client.embed([chunk.content for chunk in batch])
            for index, (chunk, vector) in enumerate(zip(batch, vectors, strict=True), start=offset):
                await session.execute(text('''
                    INSERT INTO organizations.policy_chunks
                    (policy_id,organization_id,chunk_index,content,location,token_count,embedding,embedding_model)
                    VALUES (:policy,:org,:index,:content,CAST(:location AS jsonb),:tokens,CAST(:embedding AS vector),:model)
                    ON CONFLICT(policy_id,chunk_index) DO NOTHING
                '''), {'policy': policy['id'], 'org': policy['organization_id'], 'index': index,
                       'content': chunk.content, 'location': json.dumps(chunk.location),
                       'tokens': chunk.token_count, 'embedding': json.dumps(list(vector)), 'model': MODEL})
    finally:
        if owned:
            await client.aclose()
    # Serialize replacements within a document family so only one version becomes active.
    await session.execute(text("SELECT pg_advisory_xact_lock(hashtextextended(:family,0))"),
                          {'family': f"{policy['organization_id']}:{policy['document_family_id']}"})
    newer = (await session.execute(text("""SELECT EXISTS(SELECT 1 FROM organizations.tenant_policies
        WHERE organization_id=:org AND document_family_id=:family AND active AND created_at>:created)"""),
        {'org': policy['organization_id'], 'family': policy['document_family_id'], 'created': policy['created_at']})).scalar_one()
    if not newer:
        await session.execute(text("""UPDATE organizations.tenant_policies SET active=false,updated_at=now()
            WHERE organization_id=:org AND document_family_id=:family AND active"""),
            {'org': policy['organization_id'], 'family': policy['document_family_id']})
    await session.execute(text("""UPDATE organizations.tenant_policies
        SET processing_status='ready',active=:active,embedded_chunks_count=:count,error_code=NULL,
            lease_until=NULL,updated_at=now() WHERE id=:id AND organization_id=:org"""),
        {'id': policy['id'], 'org': policy['organization_id'], 'count': len(chunks), 'active': not newer})
    return len(chunks)
